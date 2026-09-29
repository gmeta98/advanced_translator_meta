"""Regression checks using synthetic documents and mocked translation responses."""

import base64
import json
import os
import re
import unittest
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import streamlit as st
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from streamlit.testing.v1 import AppTest

import app


TRANSLATION_SETTINGS = dict(
    source_language="English", target_language="Italian",
    subject_matter="Business document", style_instruction="", glossary="",
    model="gpt-6-sol",
)


def document_bytes(document):
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def structure_without_text(element):
    element = deepcopy(element)
    for node in element.xpath(".//w:t"):
        node.text = ""
        node.attrib.pop(app.XML_SPACE, None)
    return etree.tostring(element)


class Upload(BytesIO):
    def __init__(self, name="example.docx", text="Name\tValue\nNext line"):
        document = Document()
        document.add_paragraph(text)
        super().__init__(document_bytes(document))
        self.name = name


class DocumentTests(unittest.TestCase):
    def test_tabs_breaks_runs_headers_and_footers_survive_export(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("\t Name").bold = True
        paragraph.add_run("\t\tValue\n").italic = True
        paragraph.add_run("Next line ").add_break(WD_BREAK.PAGE)
        paragraph.add_run("After page\n")
        document.sections[0].header.paragraphs[0].text = "Header\tValue"
        document.sections[0].footer.paragraphs[0].text = "Footer\nValue"
        translations = {"Name": "Nome", "Value": "Valore", "Next line": "Riga successiva",
                        "After page": "Dopo pagina", "Header": "Intestazione", "Footer": "Piè di pagina"}
        original = structure_without_text(document.element)
        progress = Mock()
        with patch.object(app, "translate_units", side_effect=lambda units, **_: [translations[u.text] for u in units]), \
                patch.object(app.st, "progress", return_value=progress):
            output, pairs = app.translate_document(document_bytes(document), user_context="", **TRANSLATION_SETTINGS)
        result = Document(BytesIO(output))
        self.assertEqual(result.paragraphs[0].text, "\t Nome\t\tValore\nRiga successiva Dopo pagina\n")
        self.assertEqual(result.sections[0].header.paragraphs[0].text, "Intestazione\tValore")
        self.assertEqual(result.sections[0].footer.paragraphs[0].text, "Piè di pagina\nValore")
        self.assertEqual(structure_without_text(result.element), original)
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertTrue(result.paragraphs[0].runs[1].italic)
        self.assertEqual(len(pairs), 8)
        progress.empty.assert_called_once()

    def test_hyperlinks_are_translated_without_changing_the_link(self):
        document = Document()
        paragraph = document.add_paragraph("Label\t")
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), "bookmark")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Linked text"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        for unit, translation in zip(app.collect_text_units(document), ["Etichetta", "Testo collegato"]):
            app.replace_text_unit(document, unit, translation)
        self.assertEqual(paragraph.text, "Etichetta\tTesto collegato")
        self.assertEqual(hyperlink.get(qn("w:anchor")), "bookmark")

    def test_nested_drawing_text_is_not_overwritten(self):
        document = Document()
        paragraph = document.add_paragraph("Main text")
        drawing = OxmlElement("w:drawing")
        inner_paragraph = OxmlElement("w:p")
        inner_run = OxmlElement("w:r")
        inner_text = OxmlElement("w:t")
        inner_text.text = "Text in drawing"
        inner_run.append(inner_text)
        inner_paragraph.append(inner_run)
        drawing.append(inner_paragraph)
        paragraph.runs[0]._r.append(drawing)
        app.replace_text_unit(document, app.collect_text_units(document)[0], "Testo principale")
        self.assertEqual(inner_text.text, "Text in drawing")

    def test_nested_tables_follow_reading_order_and_merged_cells_are_not_repeated(self):
        document = Document()
        table = document.add_table(rows=1, cols=2)
        cell = table.cell(0, 0).merge(table.cell(0, 1))
        cell.paragraphs[0].text = "Before table"
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "Inside table"
        cell.paragraphs[-1].text = "After table"
        self.assertEqual([u.text for u in app.collect_text_units(document)],
                         ["Before table", "Inside table", "After table"])

    def test_empty_replacement_leaves_source_untouched(self):
        document = Document()
        paragraph = document.add_paragraph("Original")
        unit = app.collect_text_units(document)[0]
        with self.assertRaises(ValueError):
            app.replace_text_unit(document, unit, " \t\n")
        self.assertEqual(paragraph.text, "Original")

    def test_failed_translation_clears_progress(self):
        progress = Mock()
        with patch.object(app, "translate_units", side_effect=ValueError("Empty response")), \
                patch.object(app.st, "progress", return_value=progress):
            with self.assertRaises(ValueError):
                app.translate_document(Upload().getvalue(), user_context="", **TRANSLATION_SETTINGS)
        progress.empty.assert_called_once()


class ResponseAndPasswordTests(unittest.TestCase):
    def setUp(self):
        self.units = [app.TextUnit("Body", "Original", 0)]

    def test_invalid_translations_are_rejected(self):
        invalid_items = [
            [], [{"id": 0, "text": ""}], [{"id": 0, "text": " \n\t"}],
            [{"id": 0, "text": None}], [{"id": 1, "text": "Wrong id"}],
            [{"id": False, "text": "Boolean id"}], ["Not an object"],
            [{"id": 0, "text": "One"}, {"id": 0, "text": "Duplicate"}],
            [{"id": 0, "text": "Valid"}, {"id": 9, "text": "Unexpected"}],
        ]
        for items in invalid_items:
            with self.subTest(items=items), self.assertRaises(ValueError):
                app.parse_translation_response(json.dumps({"translations": items}), self.units)

    def test_reordered_valid_responses_match_source_ids(self):
        units = self.units + [app.TextUnit("Body 2", "Next", 4)]
        response = {"translations": [{"id": 4, "text": "Secondo"}, {"id": 0, "text": "Primo"}]}
        self.assertEqual(app.parse_translation_response(json.dumps(response), units), ["Primo", "Secondo"])

    def test_empty_response_is_retried_then_validated(self):
        client = Mock()
        client.responses.create.side_effect = [
            SimpleNamespace(output_text='{"translations":[{"id":0,"text":""}]}'),
            SimpleNamespace(output_text='{"translations":[{"id":0,"text":"Tradotto"}]}'),
        ]
        with patch.object(app, "get_configured_client", return_value=client):
            result = app.translate_units(self.units, document_context="", recent_context="", **TRANSLATION_SETTINGS)
        self.assertEqual(result, ["Tradotto"])
        self.assertEqual(client.responses.create.call_count, 2)

    def test_repeated_empty_response_stops_translation(self):
        client = Mock()
        client.responses.create.return_value = SimpleNamespace(output_text='{"translations":[{"id":0,"text":""}]}')
        with patch.object(app, "get_configured_client", return_value=client), self.assertRaises(ValueError):
            app.translate_units(self.units, document_context="", recent_context="", **TRANSLATION_SETTINGS)
        self.assertEqual(client.responses.create.call_count, app.MAX_TRANSLATION_ATTEMPTS)

    def test_unicode_passwords_succeed_or_fail_without_crashing(self):
        for entered, expected in [("fjalëkalim🔑", True), ("fjalekalim🔑", False), ("", False)]:
            state = {"login_password": entered}
            with self.subTest(entered=entered), patch.object(app.st, "session_state", state), \
                    patch.object(app, "get_config_value", return_value="fjalëkalim🔑"):
                app.check_password()
            self.assertEqual(state["authenticated"], expected)
            self.assertEqual(state["login_error"], not expected)
            self.assertNotIn("login_password", state)


class SessionTests(unittest.TestCase):
    def setUp(self):
        self.uploads = [Upload()]
        self.enterContext(patch.dict(os.environ, {"APP_PASSWORD": "test-only", "OPENAI_API_KEY": ""}))
        self.enterContext(patch.object(st, "secrets", {}))
        self.enterContext(patch.object(st, "file_uploader", side_effect=lambda *a, **k: self.uploads))
        self.enterContext(patch("openai.OpenAI", side_effect=AssertionError("Tests must not call the live API")))
        self.at = AppTest.from_file(str(Path(__file__).with_name("app.py")))
        self.at.session_state["authenticated"] = True
        self.at.run()
        self.assertFalse(self.at.exception)

    def translate(self):
        next(b for b in self.at.button if b.label == "Translate document").click().run()
        self.assertFalse(self.at.exception)

    def assert_no_output(self):
        self.assertFalse(self.at.exception)
        self.assertEqual(len(self.at.get("download_button")), 0)
        self.assertEqual(len(self.at.get("iframe")), 0)
        for key in ("translated_docx", "translated_zip", "translated_pairs", "download_name", "translation_job_signature"):
            self.assertNotIn(key, self.at.session_state)

    def test_models_and_downloads_still_work_without_repeating_on_rerun(self):
        model = next(w for w in self.at.selectbox if w.label == "OpenAI model")
        self.assertEqual(model.options, ["gpt-6-sol", "gpt-6-luna"])
        self.assertEqual(model.value, "gpt-6-sol")
        self.translate()
        html = self.at.get("iframe")[0].proto.srcdoc
        encoded = re.search(r'atob\("([A-Za-z0-9+/=]+)"\)', html).group(1)
        self.assertEqual(base64.b64decode(encoded), self.at.session_state["translated_docx"])
        self.at.run()
        self.assertEqual(len(self.at.get("download_button")), 1)
        self.assertEqual(len(self.at.get("iframe")), 0)
        self.translate()
        self.assertNotEqual(self.at.get("iframe")[0].proto.srcdoc, html)

    def test_setting_changes_clear_old_downloads(self):
        changes = [
            ("selectbox", "Target language", "English"),
            ("selectbox", "Source language", "German"),
            ("selectbox", "Document type", "Business document"),
            ("selectbox", "OpenAI model", "gpt-6-luna"),
            ("text_area", "Required terminology", "Term = Translation"),
            ("text_area", "Style requirements", "Updated style"),
            ("text_area", "Document context", "Updated context"),
        ]
        for kind, label, value in changes:
            with self.subTest(label=label):
                self.translate()
                widget = next(w for w in getattr(self.at, kind) if w.label == label)
                widget.set_value(value).run()
                self.assert_no_output()

    def test_replaced_same_name_upload_and_removed_upload_clear_results(self):
        self.translate()
        self.uploads[:] = [Upload(text="Different document, same file name")]
        self.at.run()
        self.assert_no_output()
        self.translate()
        self.uploads.clear()
        self.at.run()
        self.assert_no_output()

    def test_zip_download_and_file_list_changes(self):
        self.uploads[:] = [Upload("one.docx"), Upload("two.docx")]
        self.at.run()
        self.translate()
        with zipfile.ZipFile(BytesIO(self.at.session_state["translated_zip"])) as archive:
            self.assertEqual(archive.namelist(), ["one_Albanian.docx", "two_Albanian.docx"])
        self.assertEqual(len(self.at.get("iframe")), 1)
        self.uploads.reverse()
        self.at.run()
        self.assert_no_output()

    def test_failed_retranslation_does_not_keep_old_download(self):
        def valid_response(**request):
            items = json.loads(request["input"][1]["content"].split("Text items:\n", 1)[1])
            return SimpleNamespace(output_text=json.dumps({"translations": [
                {"id": item["id"], "text": "Translated " + item["text"]} for item in items
            ]}))

        client = Mock()
        client.responses.create.side_effect = valid_response
        with patch.dict(os.environ, {"OPENAI_API_KEY": "synthetic-test-key"}), \
                patch("openai.OpenAI", return_value=client):
            self.translate()
            self.assertEqual(len(self.at.get("download_button")), 1)
            client.responses.create.side_effect = None
            client.responses.create.return_value = SimpleNamespace(output_text='{"translations":[]}')
            self.translate()
        self.assert_no_output()
        self.assertEqual(len(self.at.error), 1)

    def test_logout_clears_documents_and_signature(self):
        self.translate()
        next(b for b in self.at.button if b.label == "Log out").click().run()
        self.assert_no_output()
        self.assertNotIn("authenticated", self.at.session_state)


if __name__ == "__main__":
    unittest.main()
